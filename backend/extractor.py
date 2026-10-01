"""
extractor.py - Text extraction module for PDF and DOCX files.
Implements FR-05: extracts readable text from supported documents.
"""

import io
from pypdf import PdfReader
import docx


def extract_text_from_pdf(file_bytes: bytes) -> str:
    """
    Extract readable text from a PDF file in-memory using pypdf.
    Returns extracted text or raises ValueError if unreadable/empty.
    """
    try:
        reader = PdfReader(io.BytesIO(file_bytes))
        extracted_pages = []
        for idx, page in enumerate(reader.pages):
            text = page.extract_text()
            if text:
                extracted_pages.append(text)

        full_text = "\n".join(extracted_pages).strip()
        if not full_text:
            raise ValueError(
                "Unable to extract text from PDF. The file may be scanned, image-based, or empty."
            )
        return full_text
    except Exception as exc:
        if isinstance(exc, ValueError):
            raise exc
        raise ValueError(f"Error parsing PDF file: {str(exc)}")


def extract_text_from_docx(file_bytes: bytes) -> str:
    """
    Extract readable text from a DOCX file in-memory using python-docx.
    Returns extracted text or raises ValueError if unreadable/empty.
    """
    try:
        doc = docx.Document(io.BytesIO(file_bytes))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        # Also extract table text
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        paragraphs.append(cell.text.strip())

        full_text = "\n".join(paragraphs).strip()
        if not full_text:
            raise ValueError("Unable to extract text from DOCX. The document appears to be empty.")
        return full_text
    except Exception as exc:
        if isinstance(exc, ValueError):
            raise exc
        raise ValueError(f"Error parsing DOCX file: {str(exc)}")


def extract_text(filename: str, file_bytes: bytes) -> str:
    """Dispatches extraction according to file extension."""
    lower_name = filename.lower()
    if lower_name.endswith(".pdf"):
        return extract_text_from_pdf(file_bytes)
    elif lower_name.endswith(".docx"):
        return extract_text_from_docx(file_bytes)
    else:
        raise ValueError("Unsupported file format. Please upload a PDF or DOCX file.")
