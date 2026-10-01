"""
generate_dataset_samples.py
Generates realistic, anonymized PDF and DOCX test resumes based on
standard benchmark categories (Data Science, DevOps, Full Stack, Cybersecurity, HR).
"""

import os
import docx
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors


def create_pdf_resume(filepath: str, name: str, contact: str, summary: str, skills_text: str, experience: list[str]):
    """Builds a formatted, readable PDF resume using ReportLab."""
    doc = SimpleDocTemplate(
        filepath,
        pagesize=letter,
        rightMargin=40, leftMargin=40, topMargin=40, bottomMargin=40
    )
    styles = getSampleStyleSheet()

    # Custom styles
    title_style = ParagraphStyle(
        'CandidateName',
        parent=styles['Heading1'],
        fontSize=20,
        leading=24,
        textColor=colors.HexColor('#1e3a8a'),
        spaceAfter=4
    )
    contact_style = ParagraphStyle(
        'ContactInfo',
        parent=styles['Normal'],
        fontSize=9,
        leading=12,
        textColor=colors.HexColor('#475569'),
        spaceAfter=14
    )
    section_title = ParagraphStyle(
        'SectionHeader',
        parent=styles['Heading2'],
        fontSize=12,
        leading=16,
        textColor=colors.HexColor('#2563eb'),
        spaceBefore=10,
        spaceAfter=4
    )
    body_style = ParagraphStyle(
        'Body',
        parent=styles['Normal'],
        fontSize=10,
        leading=14,
        textColor=colors.HexColor('#1f2937'),
        spaceAfter=6
    )

    story = []
    story.append(Paragraph(name, title_style))
    story.append(Paragraph(contact, contact_style))

    story.append(Paragraph("Professional Summary", section_title))
    story.append(Paragraph(summary, body_style))

    story.append(Paragraph("Core Technical Skills", section_title))
    story.append(Paragraph(skills_text, body_style))

    story.append(Paragraph("Work Experience", section_title))
    for exp in experience:
        story.append(Paragraph(f"• {exp}", body_style))

    doc.build(story)


def generate_all_samples():
    base_dir = os.path.dirname(__file__)

    # 1. Data Scientist (PDF)
    create_pdf_resume(
        os.path.join(base_dir, "data_scientist_resume.pdf"),
        name="Dr. Elena Rostova",
        contact="Email: elena.rostova@example.com | Phone: (555) 432-8765 | Boston, MA",
        summary="Data Scientist with 4+ years developing predictive models, NLP pipelines, and deep neural networks in healthcare and finance.",
        skills_text="Python, SQL, Machine Learning, Deep Learning, PyTorch, TensorFlow, Scikit-Learn, Pandas, NumPy, NLP, RAG, Power BI",
        experience=[
            "Senior ML Engineer at BioHealth: Implemented RAG pipelines and transformer models with Hugging Face and PyTorch.",
            "Data Scientist at Alpha Analytics: Built predictive churn models using Scikit-Learn and Pandas on SQL databases.",
            "Designed automated ETL pipelines and visualized feature importance metrics using Power BI and Tableau."
        ]
    )

    # 2. DevOps & Cloud Engineer (PDF)
    create_pdf_resume(
        os.path.join(base_dir, "devops_cloud_engineer_resume.pdf"),
        name="Marcus Vance",
        contact="Email: marcus.vance@example.com | Phone: (555) 678-9012 | Austin, TX",
        summary="Senior DevOps Engineer with expertise in multi-cloud infrastructure, container orchestration, and continuous delivery.",
        skills_text="Kubernetes, Docker, AWS, Terraform, CI/CD, Linux, Python, Bash, Prometheus, Grafana, Ansible, Git",
        experience=[
            "Cloud DevOps Lead at Apex Cloud: Managed production Kubernetes (k8s) clusters on AWS across 3 regions.",
            "Automated infrastructure provisioning with Terraform and Ansible, decreasing deployment cycle time by 45%.",
            "Configured monitoring and alerting dashboards using Prometheus, Grafana, and ELK stack on Linux."
        ]
    )

    # 3. Full Stack Developer (DOCX)
    doc_fs = docx.Document()
    doc_fs.add_heading("David Martinez", 0)
    doc_fs.add_paragraph("Email: david.martinez@example.com | Phone: (555) 345-6789 | Seattle, WA")
    doc_fs.add_heading("Professional Summary", level=1)
    doc_fs.add_paragraph("Full Stack Software Engineer proficient in React, Node.js, TypeScript, and modern relational databases.")
    doc_fs.add_heading("Technical Skills", level=1)
    doc_fs.add_paragraph("Frontend: React, TypeScript, JavaScript, HTML, CSS, Next.js, Tailwind CSS\nBackend: Node.js, Express, FastAPI, Python, REST API, GraphQL\nDatabases: PostgreSQL, MongoDB, Redis\nTools: Git, Docker, Jest")
    doc_fs.add_heading("Experience", level=1)
    doc_fs.add_paragraph("Full Stack Engineer at ModernWeb (2022 - Present)\n- Engineered responsive client interfaces using Next.js, React, and Tailwind CSS.\n- Built RESTful microservices with Node.js and PostgreSQL.\n- Automated unit testing suites using Jest.")
    doc_fs.save(os.path.join(base_dir, "fullstack_developer_resume.docx"))

    # 4. Cybersecurity Analyst (PDF)
    create_pdf_resume(
        os.path.join(base_dir, "cybersecurity_analyst_resume.pdf"),
        name="Sarah Jenkins",
        contact="Email: sarah.jenkins@example.com | Phone: (555) 890-1234 | Washington, DC",
        summary="Information Security Specialist specializing in threat hunting, vulnerability management, and web application security.",
        skills_text="Network Security, Penetration Testing, OWASP, Python, Linux, Cryptography, SIEM, Ethical Hacking, Bash",
        experience=[
            "Cybersecurity Analyst at SecureDefense: Conducted web application penetration testing aligned with OWASP Top 10.",
            "Monitored enterprise SIEM alerts and automated threat triage scripts using Python and Bash on Linux servers.",
            "Administered vulnerability assessment scans and remediation reports for SOC2 audits."
        ]
    )

    # 5. HR Operations Manager (PDF - Negative Control / 0% Match)
    create_pdf_resume(
        os.path.join(base_dir, "hr_manager_resume.pdf"),
        name="Jessica Taylor",
        contact="Email: jessica.taylor@example.com | Phone: (555) 789-0123 | New York, NY",
        summary="Senior Human Resources Manager with 8+ years leading talent acquisition, employee onboarding, benefits administration, and organizational leadership.",
        skills_text="Talent Acquisition, Performance Management, Employee Relations, Payroll, Benefits Administration, Conflict Resolution, Interviewing",
        experience=[
            "Human Resources Director at Global Corp: Led full-cycle recruitment for 200+ employees annually.",
            "Spearheaded onboarding programs and employee retention initiatives across 5 regional branches.",
            "Managed organizational compliance, workplace safety, and annual performance review cycles."
        ]
    )

    # Cleanup temporary test.pdf if it exists
    test_pdf = os.path.join(base_dir, "test.pdf")
    if os.path.exists(test_pdf):
        os.remove(test_pdf)

    print("All dataset sample resumes generated successfully!")


if __name__ == "__main__":
    generate_all_samples()
